# 导入docx
import docx

# 导入openpyxl
import openpyxl

# 读取工作目录下的Excel工作簿"员工花名册.xlsx"，并赋值给变量wb
wb = openpyxl.load_workbook("员工花名册.xlsx")

# 读取工作表"汇总"，并赋值给变量ws
ws = wb["汇总"]

# 遍历工作表的每一行
for row in ws.rows:

    # TODO 读取每一行的第一个单元格值，并赋值给变量name
    name = row[0].value
    
    if name == "姓名":
        continue
    else:
        
        # TODO 读取工作目录下的Word文档"八周年邀请函模版.docx"，并赋值给变量doc
        doc = docx.Document("八周年邀请函模版.docx")

        # TODO 使用add_text()函数，在doc的第4个段落第1个样式块末尾，添加变量name的文本内容
        doc.paragraphs[3].runs[0].add_text(name)

        # TODO 储存Word文档至文件夹"邀请函"下，并命名为"八周年邀请函_致xxx.docx"
        doc.save(f"邀请函\\八周年邀请函_致{name}.docx")

# 如果操作成功，输出"success"
print("success")